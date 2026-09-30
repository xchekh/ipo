"""
Сборка региональных ипотечных таблиц с нуля из сырой выгрузки.

Вход:  «Статистические ряды_регионы.xlsx» (листы 01_02_01, 01_02_03, 01_02_05).
Выход: новый xlsx только с готовыми листами «кол-во», «объем», «ДДУ»,
       «01_02_05» (характеристики кредитов) в формате таблицы Дани, с уникальными тикерами.

Правила (как у Дани):
- кол-во / объем — строки 01_02_01 с размерностью «Всего, шт.» / «Всего, млн руб.»,
  без «Все программы» и «Ипотека в отдельных регионах».
  Тикер = <тикер региона> + название программы («Город Москва Льготная ипотека») —
  формулой Дани =$B$<строка региона>&A; тикер региона стоит значением в строке региона.
- ДДУ — строки 01_02_03 «Покупка по ДДУ, шт./млн руб.» по всем программам.
  Ключ = <тикер региона><программа><строка> («Город МоскваВсе программы   Покупка по ДДУ, шт.») —
  формулой Дани =СМЕЩ(D;-A;0)&СМЕЩ(D;-B;-1)&C: A — позиция строки в блоке региона
  (19 строк), B — 1/2 для шт./млн руб.; тикер региона стоит значением в строке региона.
- 01_02_05 (характеристики кредитов) — все строки выгрузки 01_02_05 в раскладке Дани:
  B — тикер региона, C — ключ, данные с D. Ключ и числа стоят только в строке региона,
  формулами Дани: C = B&A(+1)&A(+5) («Город МоскваВсе программы   Средний срок кредита, мес.»),
  данные = ссылки на строку «Средний срок кредита, мес.» по всем программам.
- Тикер региона — как у Дани (REGION_TICKERS); для остальных регионов — само название.

Запуск:  python build_regional_tables.py "Статистические ряды_регионы.xlsx" [выходной.xlsx]
Нужен:   pip install openpyxl
"""
import re
import sys
import zipfile
from copy import copy
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import get_column_letter

# Тикеры регионов из таблицы Дани, отличающиеся от названия региона в выгрузке.
# Все прочие регионы (включая ФО, у которых у Дани стоит 0) — тикер = название.
REGION_TICKERS = {
    "Российская Федерация": "Всего РФ",
    "г. Москва": "Город Москва",
    "г. Санкт-Петербург": "Город Санкт-Петербург",
    "г. Севастополь": "Город Севастополь",
    "Республика Татарстан (Татарстан)": "Республика Татарстан",
    "Чувашская Республика - Чувашия": "Чувашская Республика",
    "Республика Адыгея (Адыгея)": "Республика Адыгея",
    "Республика Северная Осетия - Алания": "Республика Северная Осетия",
}

EXCLUDED_PROGRAMS = {"Все программы", "Ипотека в отдельных регионах"}
DDU_RE = re.compile(r"^\s*Покупка по ДДУ,\s*(шт\.|млн руб\.)$")
FIRST_DATA_ROW = 5
HEADER_ROW = 4
DDU_BLOCK = 19  # строка региона + 6 программ × (заголовок + шт. + млн руб.)
CHAR_SHEET = "01_02_05"  # имя готового листа характеристик — как у Дани
CHAR_TMP_TITLE = "Характеристики кредитов"  # временное имя, пока сырой 01_02_05 ещё в книге
CHAR_KEY_PROGRAM = "Все программы"
CHAR_KEY_ROW = "Средний срок кредита, мес."


def region_ticker(name):
    return REGION_TICKERS.get(name, name)


def region_rows(ws):
    """Строки-заголовки регионов: за ними сразу идёт «Все программы»."""
    last = ws.max_row
    labels = {r: ws.cell(r, 1).value for r in range(FIRST_DATA_ROW, last + 1)}
    return {r for r in range(FIRST_DATA_ROW, last)
            if labels[r + 1] == "Все программы" and labels[r] != "Все программы"}


def col_widths(ws):
    widths = {}
    for dim in ws.column_dimensions.values():
        if dim.width and dim.min and dim.max:
            for c in range(dim.min, dim.max + 1):
                widths[c] = dim.width
    return widths


class SheetWriter:
    """Копирует строки исходного листа в новый, со сдвигом колонок на `shift`."""

    def __init__(self, wb, title, src, shift, top_rows=(1, 2, 3, 4)):
        self.ws = wb.create_sheet(title)
        self.src = src
        self.shift = shift
        self.ncols = src.max_column
        self.row = 0
        for w_col, width in col_widths(src).items():
            self.ws.column_dimensions[get_column_letter(w_col + shift)].width = width
        for r in top_rows:
            self.copy_row(r)

    def copy_row(self, src_row):
        self.row += 1
        for c in range(1, self.ncols + 1):
            s = self.src.cell(src_row, c)
            d = self.ws.cell(self.row, c + self.shift)
            d.value = s.value
            if s.has_style:
                d._style = copy(s._style)
        return self.row

    def set(self, col, value, style_from=None):
        cell = self.ws.cell(self.row, col, value)
        if style_from is not None:
            cell._style = copy(self.ws.cell(self.row, style_from)._style)
        return cell

    def finish(self, freeze, filter_first_col):
        last = get_column_letter(self.ncols + self.shift)
        self.ws.freeze_panes = freeze
        self.ws.auto_filter.ref = f"{filter_first_col}{HEADER_ROW}:{last}{self.row}"


def build_count_volume(wb, src, title, dim_label):
    w = SheetWriter(wb, title, src, shift=1)
    fix_first_col(w, range(1, HEADER_ROW + 1))
    w.ws.column_dimensions["A"].width = col_widths(src).get(1, 24.7)
    w.ws.column_dimensions["B"].width = 24.7
    w.ws.cell(HEADER_ROW, 2, "Тикер")._style = copy(w.ws.cell(HEADER_ROW, 1)._style)
    # служебная строка 3 с номерами колонок для ВПР (как у Дани)
    w.ws.cell(3, 2, 1)
    for c in range(3, src.max_column + 2):
        w.ws.cell(3, c, c)

    regions = region_rows(src)
    prefix = reg_row = None
    tickers = []
    cached = {}  # результаты формул — чтобы тикеры читались и без пересчёта в Excel
    for r in range(FIRST_DATA_ROW, src.max_row + 1):
        label = src.cell(r, 1).value
        if r in regions:
            prefix = region_ticker(label)
            w.copy_row(r)
            fix_first_col(w, [w.row])
            w.set(2, prefix, style_from=1)
            reg_row = w.row
            tickers.append(prefix)
            continue
        if src.cell(r, 2).value != dim_label or not isinstance(label, str):
            continue
        if label.strip() in EXCLUDED_PROGRAMS:
            continue
        w.copy_row(r)
        fix_first_col(w, [w.row])
        # формула Дани: тикер региона из строки региона + название программы
        w.set(2, f"=$B${reg_row}&A{w.row}")
        tickers.append(prefix + label)
        cached[f"B{w.row}"] = prefix + label
    w.finish("D5", "A")
    return tickers, cached


def fix_first_col(w, rows):
    """copy_row сдвигает всю строку вправо; вернуть колонку подписей в A."""
    for rr in rows:
        b = w.ws.cell(rr, 2)
        a = w.ws.cell(rr, 1)
        a.value = b.value
        a._style = copy(b._style)
        b.value = None


def build_ddu(wb, src):
    w = SheetWriter(wb, "ДДУ", src, shift=3)
    # колонка подписей: из A выгрузки в C, данные с E; D — ключ
    for rr in range(1, w.row + 1):
        move(w.ws, rr, 4, 3)
    w.ws.column_dimensions["A"].width = 8.1
    w.ws.column_dimensions["B"].width = 8.1
    w.ws.column_dimensions["C"].width = col_widths(src).get(1, 47.7)
    w.ws.column_dimensions["D"].width = 46.3
    for c in range(4, src.max_column + 4):
        w.ws.cell(3, c, c - 3)

    regions = region_rows(src)
    prefix = program = None
    keys = []
    cached = {}  # результаты формул A/B/D — чтобы ключи читались и без пересчёта в Excel
    block_start = None

    def check_block():
        # формулы Дани в A/D рассчитаны на блок ровно из DDU_BLOCK строк
        if block_start is not None and w.row - block_start + 1 != DDU_BLOCK:
            raise SystemExit(
                f"ДДУ: у региона «{w.ws.cell(block_start, 3).value}» {w.row - block_start + 1} строк "
                f"вместо {DDU_BLOCK} — в выгрузке изменился набор программ, формулы Дани не подойдут")

    for r in range(FIRST_DATA_ROW, src.max_row + 1):
        label = src.cell(r, 1).value
        if not isinstance(label, str):
            continue
        if r in regions:
            check_block()
            block_start = w.row + 1
            prefix, program, key = region_ticker(label), None, region_ticker(label)
        elif src.cell(r, 2).value is None and not label.startswith(" "):
            program, key = label, None
        elif DDU_RE.match(label):
            key = "formula"
            keys.append(prefix + program + label)
        else:
            continue
        w.copy_row(r)
        move(w.ws, w.row, 4, 3)
        rr = w.row
        if key == "formula":
            key = f"=OFFSET(D{rr},-(A{rr}),0)&OFFSET(D{rr},-(B{rr}),-1)&C{rr}"
            cached[f"D{rr}"] = keys[-1]
        w.set(4, key, style_from=3 if key is None else None)
        w.ws.cell(rr, 1, 0 if rr == FIRST_DATA_ROW else f"=IF(A{rr-1}=18,0,A{rr-1}+1)")
        w.ws.cell(rr, 2, f"=IF(C{rr}=$C$7,1,IF(C{rr}=$C$8,2,0))")
        cached[f"A{rr}"] = rr - block_start
        m = DDU_RE.match(label)
        cached[f"B{rr}"] = {"шт.": 1, "млн руб.": 2}[m.group(1)] if m else 0
    check_block()
    w.finish("D5", "C")
    return keys, cached


def move(ws, row, from_col, to_col):
    s, d = ws.cell(row, from_col), ws.cell(row, to_col)
    d.value, d._style = s.value, copy(s._style)
    s.value = None


def build_characteristics(wb, src):
    # раскладка и формулы как на 01_02_05 у Дани: A — подписи, B — тикер региона, C — ключ, данные с D;
    # ключ и числа стоят в строке региона и дублируют «Средний срок кредита» по всем программам
    w = SheetWriter(wb, CHAR_TMP_TITLE, src, shift=2)
    for rr in range(1, w.row + 1):
        move(w.ws, rr, 3, 1)
    w.ws.column_dimensions["A"].width = col_widths(src).get(1, 53.7)
    w.ws.column_dimensions["B"].width = 24.7
    w.ws.column_dimensions["C"].width = 46.3
    w.ws.cell(HEADER_ROW, 2, "Тикер")._style = copy(w.ws.cell(HEADER_ROW, 1)._style)
    w.ws.cell(HEADER_ROW, 3, "Ключ")._style = copy(w.ws.cell(HEADER_ROW, 1)._style)
    # служебная строка 3 с номерами колонок для ВПР (как у Дани): C = 1, D = 2, ...
    for c in range(3, src.max_column + 3):
        w.ws.cell(3, c, c - 2)

    regions = region_rows(src)
    keys = []
    cached = {}  # результаты формул — чтобы ключи и числа читались и без пересчёта в Excel
    reg_row = reg_name = None  # строка региона, ещё не получившая ключ

    def check_region():
        if reg_row is not None:
            raise SystemExit(
                f"Характеристики: у региона «{reg_name}» нет строки «{CHAR_KEY_PROGRAM}» / «{CHAR_KEY_ROW}» "
                f"— в выгрузке изменилась структура, формулы Дани не подойдут")

    program = None
    for r in range(FIRST_DATA_ROW, src.max_row + 1):
        label = src.cell(r, 1).value
        if not isinstance(label, str):
            continue
        w.copy_row(r)
        move(w.ws, w.row, 3, 1)
        if r in regions:
            check_region()
            reg_row, reg_name, program = w.row, label, None
            w.set(2, region_ticker(label), style_from=1)
        elif not label.startswith(" "):
            program = label
        elif reg_row is not None and program == CHAR_KEY_PROGRAM and label.strip() == CHAR_KEY_ROW:
            # у Дани: C = B&A(+1)&A(+5), данные = ссылки на строку «Средний срок кредита»
            rr = w.row
            w.ws.cell(reg_row, 3, f"=B{reg_row}&A{reg_row + 1}&A{rr}")._style = copy(w.ws.cell(reg_row, 1)._style)
            key = region_ticker(reg_name) + w.ws.cell(reg_row + 1, 1).value + label
            keys.append(key)
            cached[f"C{reg_row}"] = key
            for c in range(4, w.ncols + w.shift + 1):
                col = get_column_letter(c)
                value = w.ws.cell(rr, c).value
                cell = w.ws.cell(reg_row, c, f"={col}{rr}")
                cell.number_format = w.ws.cell(rr, c).number_format
                cached[f"{col}{reg_row}"] = 0 if value is None else value
            reg_row = None
    check_region()
    w.finish("D5", "A")
    return keys, cached


def ticker_values(ws, col):
    return [ws.cell(r, col).value for r in range(FIRST_DATA_ROW, ws.max_row + 1)]


def check_unique(values):
    seen, dups = set(), set()
    for v in values:
        if v in (None, ""):
            continue
        (dups if v in seen else seen).add(v)
    return len(seen), dups


def add_cached_values(path, sheet_name, values):
    """openpyxl пишет формулы без результата; дописываем результат в XML листа, как это делает Excel.
    Без этого ключи пусты для ВПР из закрытой книги и для pandas, пока файл не пересохранят в Excel."""
    with zipfile.ZipFile(path) as z:
        items = [(info, z.read(info.filename)) for info in z.infolist()]
    files = {info.filename: data for info, data in items}
    wb_xml = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rid = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(escape(sheet_name)), wb_xml).group(1)
    rel = next(r for r in re.findall(r"<Relationship [^>]*>", rels) if f'Id="{rid}"' in r)
    target = re.search(r'Target="([^"]+)"', rel).group(1).lstrip("/")
    sheet_path = target if target.startswith("xl/") else "xl/" + target

    def fill(m):
        ref, attrs, formula = m.groups()
        if ref not in values:
            return m.group(0)
        v = values[ref]
        if isinstance(v, str):
            return f'<c r="{ref}"{attrs} t="str"><f>{formula}</f><v>{escape(v)}</v></c>'
        return f'<c r="{ref}"{attrs}><f>{formula}</f><v>{v}</v></c>'

    xml = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', fill,
                 files[sheet_path].decode("utf-8"))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            z.writestr(info, xml.encode("utf-8") if info.filename == sheet_path else data)


def build(raw_path, out_path):
    wb = openpyxl.load_workbook(raw_path)
    for name in ("01_02_01", "01_02_03", "01_02_05"):
        if name not in wb.sheetnames:
            raise SystemExit(f"В выгрузке нет листа {name} — это точно «Статистические ряды_регионы»?")
    raw_names = list(wb.sheetnames)
    print("Собираю «кол-во»...")
    count_tickers, count_cached = build_count_volume(wb, wb["01_02_01"], "кол-во", "Всего, шт.")
    print("Собираю «объем»...")
    volume_tickers, volume_cached = build_count_volume(wb, wb["01_02_01"], "объем", "Всего, млн руб.")
    print("Собираю «ДДУ»...")
    ddu_keys, ddu_cached = build_ddu(wb, wb["01_02_03"])
    print(f"Собираю «{CHAR_SHEET}» (характеристики кредитов)...")
    char_keys, char_cached = build_characteristics(wb, wb["01_02_05"])
    for name in raw_names:
        del wb[name]
    wb[CHAR_TMP_TITLE].title = CHAR_SHEET
    wb.active = 0

    ok = True
    ddu_regions = [v for v in ticker_values(wb["ДДУ"], 4) if v and not str(v).startswith("=")]
    tickers = {  # тикеры и ключи — формулы, проверяем их расчётные значения
        "кол-во": count_tickers,
        "объем": volume_tickers,
        "ДДУ": ddu_regions + ddu_keys,
        CHAR_SHEET: char_keys,
    }
    for name, values in tickers.items():
        n, dups = check_unique(values)
        print(f"  {name}: {wb[name].max_row} строк, {n} тикеров, дублей: {len(dups)}")
        if dups:
            ok = False
            print("    дубли:", sorted(dups)[:10])
    wb.save(out_path)
    add_cached_values(out_path, "кол-во", count_cached)
    add_cached_values(out_path, "объем", volume_cached)
    add_cached_values(out_path, "ДДУ", ddu_cached)
    add_cached_values(out_path, CHAR_SHEET, char_cached)
    print(f"Готово: {out_path}" + ("" if ok else "  (ВНИМАНИЕ: есть дубли тикеров)"))
    return ok


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    raw = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "Региональные таблицы (кол-во, объем, ДДУ, характеристики).xlsx"
    sys.exit(0 if build(raw, out) else 1)
